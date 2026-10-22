import csv
import io
import json
import os
import threading
import time
import zipfile
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from datetime import datetime, timezone
import cv2
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from ..config import settings as cfg
from ..db import get_db
from ..deps import current_user, require_roles
from ..measure import haversine_m
from ..models import Measurement, Report, Setting, Trip, User
from ..services import depth_model, enrich, get_budget
from .reports import audit
from .workflow import check_access

router = APIRouter(prefix="/api", tags=["insights"])
admin = require_roles("admin")
staff = require_roles("officer", "admin")

# ---------------- budget settings ----------------
class Budget(BaseModel):
    currency: str = Field(min_length=1, max_length=4)
    low: float = Field(ge=0, le=10_000_000)
    medium: float = Field(ge=0, le=10_000_000)
    high: float = Field(ge=0, le=10_000_000)
    per_m2: float = Field(ge=0, le=10_000_000)

@router.get("/settings/budget")
def read_budget(_: User = Depends(admin), db: Session = Depends(get_db)):
    return get_budget(db)

@router.put("/settings/budget")
def save_budget(body: Budget, me: User = Depends(admin), db: Session = Depends(get_db)):
    row = db.get(Setting, "budget")
    if row:
        row.value = body.model_dump()
    else:
        db.add(Setting(key="budget", value=body.model_dump()))
    audit(db, me, "settings.budget", "updated")
    db.commit()
    return get_budget(db)

# ---------------- on-site measurements (ground truth) ----------------
class MeasureIn(BaseModel):
    width_cm: float | None = Field(None, ge=1, le=500)
    length_cm: float | None = Field(None, ge=1, le=800)
    depth_cm: float | None = Field(None, ge=0.1, le=60)

@router.post("/reports/{report_id}/measurements", status_code=201)
def add_measurement(report_id: int, body: MeasureIn, user: User = Depends(staff), db: Session = Depends(get_db)):
    r = db.get(Report, report_id)
    if not r:
        raise HTTPException(404, "Report not found")
    check_access(user, r)
    if body.width_cm is None and body.length_cm is None and body.depth_cm is None:
        raise HTTPException(422, "Enter at least one measurement")
    m = Measurement(report_id=r.id, user_id=user.id, **body.model_dump())
    db.add(m)
    audit(db, user, "measurement.add", str(r.id))
    db.commit()
    return {"id": m.id}

@router.get("/validation")
def validation(_: User = Depends(admin), db: Session = Depends(get_db)):
    """How far the camera-based estimates are from tape measurements. Nothing here is simulated."""
    latest = {}
    for m, r in db.execute(select(Measurement, Report).join(Report, Report.id == Measurement.report_id).order_by(Measurement.id)).all():
        latest[r.id] = (m, r)
    rows, w_err, l_err, w_pct, l_pct = [], [], [], [], []
    for m, r in latest.values():
        row = {"report_id": r.id, "est_width_cm": r.est_width_cm, "meas_width_cm": m.width_cm,
               "est_length_cm": r.est_length_cm, "meas_length_cm": m.length_cm, "meas_depth_cm": m.depth_cm}
        if m.width_cm and r.est_width_cm:
            w_err.append(abs(r.est_width_cm - m.width_cm)); w_pct.append(100 * w_err[-1] / m.width_cm)
        if m.length_cm and r.est_length_cm:
            l_err.append(abs(r.est_length_cm - m.length_cm)); l_pct.append(100 * l_err[-1] / m.length_cm)
        rows.append(row)
    avg = lambda xs: round(sum(xs) / len(xs), 1) if xs else None
    return {
        "measured_reports": len(latest),
        "width": {"n": len(w_err), "mae_cm": avg(w_err), "mape_pct": avg(w_pct)},
        "length": {"n": len(l_err), "mae_cm": avg(l_err), "mape_pct": avg(l_pct)},
        "depth_model": depth_model(db),
        "rows": rows[-100:],
    }

# ---------------- AI review + training export ----------------
class Review(BaseModel):
    verdict: str | None = None

@router.post("/reports/{report_id}/review")
def review_report(report_id: int, body: Review, me: User = Depends(admin), db: Session = Depends(get_db)):
    if body.verdict not in (None, "correct", "false_positive"):
        raise HTTPException(422, "verdict must be 'correct', 'false_positive' or null")
    r = db.get(Report, report_id)
    if not r:
        raise HTTPException(404, "Report not found")
    r.review = body.verdict
    audit(db, me, "report.review", f"{r.id}:{body.verdict}")
    db.commit()
    return {"id": r.id, "review": r.review}

@router.get("/training/export")
def training_export(_: User = Depends(admin), db: Session = Depends(get_db)):
    """YOLO-format dataset from admin-reviewed reports: 'correct' boxes become labels, frames with only
    false alarms become negative (empty-label) samples. Missed potholes cannot be captured here."""
    rs = db.scalars(select(Report).where(Report.review.is_not(None), Report.clean_image_path.is_not(None)).order_by(Report.id).limit(5000)).all()
    groups: dict[str, list[Report]] = {}
    for r in rs:
        groups.setdefault(r.clean_image_path, []).append(r)
    buf, frames, pos, neg = io.BytesIO(), 0, 0, 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, items in groups.items():
            path = os.path.join(cfg.evidence_dir, os.path.basename(name))
            img = cv2.imread(path)
            if img is None:
                continue
            h, w = img.shape[:2]
            lines = []
            for r in items:
                if r.review != "correct":
                    continue
                for d in r.detections:
                    x1, x2 = max(0, d["x1"]), min(w, d["x2"]); y1, y2 = max(0, d["y1"]), min(h, d["y2"])
                    lines.append(f"0 {(x1 + x2) / 2 / w:.6f} {(y1 + y2) / 2 / h:.6f} {(x2 - x1) / w:.6f} {(y2 - y1) / h:.6f}")
            z.write(path, f"images/{os.path.basename(name)}")
            z.writestr(f"labels/{os.path.splitext(os.path.basename(name))[0]}.txt", "\n".join(lines))
            frames += 1; pos += len(lines); neg += 0 if lines else 1
        z.writestr("data.yaml", "path: .\ntrain: images\nval: images\nnames:\n  0: pothole\n")
        z.writestr("README.txt", "Exported from SmartRoad AI. Split into train/val/test before training, or merge with other datasets using ml/merge_datasets.py.\n")
    return Response(buf.getvalue(), media_type="application/zip", headers={
        "Content-Disposition": 'attachment; filename="smartroad_training_data.zip"',
        "X-Frames": str(frames), "X-Positive": str(pos), "X-Negative": str(neg)})

# ---------------- monthly report export ----------------
HEADERS = ["Report ID", "Date (UTC)", "Time (UTC)", "Latitude", "Longitude", "GPS accuracy (m)", "Status", "Severity",
           "Danger %", "Danger level", "AI confidence %", "Width (cm)", "Length (cm)", "Depth (cm)", "Size source",
           "Depth source", "Area (m2)", "Estimated repair cost", "Reports merged", "Duplicate of", "Source",
           "Reported by", "Assigned to", "AI review", "Map link"]

def _row(r: Report, it: dict, users: dict) -> list:
    t = r.created_at.astimezone(timezone.utc)
    return [r.id, t.date().isoformat(), t.strftime("%H:%M:%S"), round(it["latitude"], 6), round(it["longitude"], 6),
            round(r.gps_accuracy_m) if r.gps_accuracy_m is not None else None, r.status, r.severity,
            it["danger_pct"], it["danger_level"], round(r.max_confidence * 100, 1), it["width_cm"], it["length_cm"],
            round(it["depth_cm"], 1) if it["depth_cm"] is not None else None, it["size_source"], it["depth_source"],
            round(it["area_m2"], 3) if it["area_m2"] is not None else None, it["est_cost"], r.confirmations,
            r.duplicate_of, r.source, users.get(r.user_id, ""), users.get(r.assigned_to, "") if r.assigned_to else "",
            r.review or "", f"https://www.google.com/maps?q={it['latitude']},{it['longitude']}"]

@router.get("/export/monthly")
def export_monthly(month: str = Query(..., pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
                   format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
                   _: User = Depends(admin), db: Session = Depends(get_db)):
    y, mo = (int(x) for x in month.split("-"))
    start = datetime(y, mo, 1, tzinfo=timezone.utc)
    end = datetime(y + (mo == 12), mo % 12 + 1, 1, tzinfo=timezone.utc)
    rows = db.execute(select(Report, func.ST_X(Report.location), func.ST_Y(Report.location))
                      .where(Report.created_at >= start, Report.created_at < end).order_by(Report.created_at)).all()
    items = enrich(db, rows)
    users = {u.id: u.email for u in db.scalars(select(User)).all()}
    table = [_row(r, it, users) for (r, _x, _y), it in zip(rows, items)]
    fname = f"smartroad_report_{month}"

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(HEADERS); w.writerows(table)
        return Response(("\ufeff" + buf.getvalue()).encode("utf-8"), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl is not installed in the backend image")
    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"
    ws.append(HEADERS)
    for row in table:
        ws.append(row)
    head_fill = PatternFill("solid", fgColor="0B1F4B")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(11, min(34, len(h) + 4))
    ws.column_dimensions[get_column_letter(len(HEADERS))].width = 14
    ws.column_dimensions[get_column_letter(HEADERS.index("Reported by") + 1)].width = 28
    ws.column_dimensions[get_column_letter(HEADERS.index("Assigned to") + 1)].width = 28
    cost_col = HEADERS.index("Estimated repair cost") + 1
    for r_i in range(2, len(table) + 2):
        ws.cell(r_i, cost_col).number_format = "#,##0"
        link = ws.cell(r_i, len(HEADERS))
        link.hyperlink = link.value; link.value = "Open map"; link.font = Font(color="2563EB", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    prim = [(r, it) for (r, _x, _y), it in zip(rows, items) if r.duplicate_of is None]
    cur = get_budget(db)["currency"]
    openp = [(r, it) for r, it in prim if r.status not in ("verified", "rejected")]
    sm = wb.create_sheet("Summary")
    lines = [("SmartRoad AI monthly report", month), ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
             ("", ""), ("Total reports (incl. repeat sightings)", len(table)), ("Unique hazards", len(prim)),
             ("Repairs verified", sum(1 for r, _ in prim if r.status == "verified")),
             ("Average danger %", round(sum(it["danger_pct"] for _, it in prim) / len(prim)) if prim else None),
             (f"Estimated cost, all hazards ({cur})", sum(it["est_cost"] for _, it in prim)),
             (f"Estimated cost, still open ({cur})", sum(it["est_cost"] for _, it in openp)), ("", "")]
    for title, key in (("By status", lambda r, it: r.status), ("By severity", lambda r, it: r.severity), ("By danger level", lambda r, it: it["danger_level"])):
        lines.append((title, ""))
        for v in sorted({key(r, it) for r, it in prim}):
            lines.append((f"  {v}", sum(1 for r, it in prim if key(r, it) == v)))
        lines.append(("", ""))
    lines.append(("Note", "Costs use the rates in Settings. Size and danger are estimates; see the validation card on the dashboard."))
    for a, b in lines:
        sm.append([a, b])
    sm["A1"].font = Font(bold=True, size=14)
    sm.column_dimensions["A"].width = 42; sm.column_dimensions["B"].width = 26
    out = io.BytesIO(); wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})

# ---------------- trips ----------------
class TripPoints(BaseModel):
    points: list[list[float]] = Field(default_factory=list, max_length=6000)  # [lng, lat, timestamp_ms]

def _trip_out(db: Session, t: Trip) -> dict:
    n = db.scalar(select(func.count()).select_from(Report).where(Report.trip_id == t.id)) or 0
    return {"id": t.id, "status": t.status, "started_at": t.started_at.isoformat(),
            "ended_at": t.ended_at.isoformat() if t.ended_at else None,
            "distance_m": t.distance_m or 0, "duration_s": t.duration_s or 0, "reports": n}

@router.post("/trips", status_code=201)
def start_trip(user: User = Depends(current_user), db: Session = Depends(get_db)):
    for old in db.scalars(select(Trip).where(Trip.user_id == user.id, Trip.status == "active")).all():
        old.status = "abandoned"
    t = Trip(user_id=user.id, status="active")
    db.add(t)
    audit(db, user, "trip.start", "new")
    db.commit()
    db.refresh(t)
    return _trip_out(db, t)

@router.post("/trips/{trip_id}/finish")
def finish_trip(trip_id: int, body: TripPoints, user: User = Depends(current_user), db: Session = Depends(get_db)):
    t = db.get(Trip, trip_id)
    if not t or t.user_id != user.id:
        raise HTTPException(404, "Trip not found")
    if t.status != "active":
        raise HTTPException(409, "Trip already finished")
    pts = [p for p in body.points if len(p) >= 2 and -180 <= p[0] <= 180 and -90 <= p[1] <= 90]
    dist = sum(haversine_m(pts[i][1], pts[i][0], pts[i + 1][1], pts[i + 1][0]) for i in range(len(pts) - 1))
    now = datetime.now(timezone.utc)
    t.status, t.ended_at, t.distance_m = "finished", now, dist
    t.duration_s = int((now - t.started_at).total_seconds())
    if len(pts) >= 2:
        wkt = "LINESTRING(" + ",".join(f"{p[0]} {p[1]}" for p in pts) + ")"
        t.path = func.ST_SetSRID(func.ST_GeomFromText(wkt), 4326)
    audit(db, user, "trip.finish", str(t.id))
    db.commit()
    db.refresh(t)
    return _trip_out(db, t)

@router.get("/trips")
def list_trips(user: User = Depends(current_user), db: Session = Depends(get_db)):
    ts = db.scalars(select(Trip).where(Trip.user_id == user.id, Trip.status == "finished").order_by(Trip.started_at.desc()).limit(50)).all()
    return [_trip_out(db, t) for t in ts]

@router.get("/trips/{trip_id}")
def trip_detail(trip_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    t = db.get(Trip, trip_id)
    if not t or (t.user_id != user.id and user.role != "admin"):
        raise HTTPException(404, "Trip not found")
    gj = db.scalar(select(func.ST_AsGeoJSON(Trip.path)).where(Trip.id == t.id))
    path = json.loads(gj)["coordinates"] if gj else []
    rows = db.execute(select(Report, func.ST_X(Report.location), func.ST_Y(Report.location)).where(Report.trip_id == t.id).order_by(Report.id)).all()
    return {**_trip_out(db, t), "path": path, "reports_list": enrich(db, rows)}


# ---------------- place names (reverse geocoding via OpenStreetMap Nominatim) ----------------
_place_cache: dict[tuple[float, float], dict] = {}
_place_lock = threading.Lock()
_place_state = {"last": 0.0, "fail_until": 0.0}
ROAD_TYPES = {"motorway": "Motorway", "trunk": "Highway", "primary": "Primary road", "secondary": "Secondary road",
              "tertiary": "Tertiary road", "residential": "Residential road", "service": "Service road",
              "unclassified": "Local road", "living_street": "Living street", "track": "Track"}

@router.get("/place")
def place(lat: float = Query(..., ge=-90, le=90), lng: float = Query(..., ge=-180, le=180), _: User = Depends(current_user)):
    """Street/area name and road class for a coordinate. Results are cached and calls are spaced out
    (Nominatim's usage policy allows about one request per second). This sends the coordinate to a third party."""
    key = (round(lat, 4), round(lng, 4))
    empty = {"name": None, "area": None, "road_type": None}
    if key in _place_cache:
        return _place_cache[key]
    with _place_lock:
        if key in _place_cache:
            return _place_cache[key]
        if time.time() < _place_state["fail_until"]:
            return empty
        wait = 1.1 - (time.time() - _place_state["last"])
        if wait > 0:
            time.sleep(wait)
        _place_state["last"] = time.time()
        try:
            url = "https://nominatim.openstreetmap.org/reverse?" + urlencode({"format": "jsonv2", "lat": lat, "lon": lng, "zoom": 17, "addressdetails": 1})
            req = Request(url, headers={"User-Agent": "SmartRoadAI/1.0 (road hazard reporting)", "Accept-Language": "en"})
            with urlopen(req, timeout=6) as resp:
                data = json.load(resp)
        except Exception:
            _place_state["fail_until"] = time.time() + 60  # back off after a failure
            return empty
        a = data.get("address", {})
        road = a.get("road")
        sub = a.get("suburb") or a.get("neighbourhood") or a.get("quarter")
        city = a.get("city") or a.get("town") or a.get("village") or a.get("county") or a.get("state_district")
        result = {
            "name": road or (f"Near {sub}" if sub else None),
            "area": ", ".join(x for x in [sub if road else None, city, a.get("state")] if x) or None,
            "road_type": ROAD_TYPES.get(data.get("type")) if data.get("category") == "highway" else None,
        }
        if len(_place_cache) > 5000:
            _place_cache.clear()
        _place_cache[key] = result
        return result
